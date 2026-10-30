"""Fixtures build real files on the fly, so tests need no binary blobs in git."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from docparse.infrastructure.config.settings import Settings
from docparse.infrastructure.di.container import Container, build_container


@pytest.fixture
def settings() -> Settings:
    return Settings(_env_file=None)


@pytest.fixture
def container(settings: Settings) -> Container:
    return build_container(settings)


@pytest.fixture
def text_pdf(tmp_path: Path) -> Path:
    """Two pages: a heading, paragraphs, a ruled table; page 2 is plain text."""
    path = tmp_path / "report.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, h - 72, "Quarterly Report")
    c.setFont("Helvetica", 11)
    y = h - 110
    for line in [
        "Revenue grew in every region this quarter.",
        "Costs stayed flat thanks to automation.",
    ]:
        c.drawString(72, y, line)
        y -= 16
    y -= 20
    c.drawString(72, y, "Margins improved for the third quarter in a row.")
    top, left, col_w, row_h = y - 40, 72, 120, 20
    rows = [["Region", "Revenue", "Growth"], ["EMEA", "120", "5%"], ["APAC", "80", "9%"]]
    for r, row in enumerate(rows):
        for ci, cell in enumerate(row):
            x, yy = left + ci * col_w, top - r * row_h
            c.rect(x, yy - row_h, col_w, row_h)
            c.drawString(x + 4, yy - row_h + 6, cell)
    c.showPage()
    c.setFont("Helvetica", 11)
    c.drawString(72, h - 72, "Page two has only prose and no table at all.")
    c.save()
    return path


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    """One page that is a raster image of text: no text layer, so OCR is required."""
    img = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default()
    draw.text((100, 150), "SCANNED INVOICE 2024", fill="black", font=font)
    draw.text((100, 260), "Total amount 1500 EUR", fill="black", font=font)
    path = tmp_path / "scan.pdf"
    img.save(path, "PDF", resolution=150)
    return path


@pytest.fixture
def csv_file(tmp_path: Path) -> Path:
    path = tmp_path / "data.csv"
    with path.open("w", newline="") as fh:
        csv.writer(fh, delimiter=";").writerows([["name", "qty"], ["bolt", "10"], ["nut", "25"]])
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    from docx import Document

    d = Document()
    d.add_heading("Contract", level=1)
    d.add_paragraph("This agreement is made between A and B.")
    d.add_heading("Terms", level=2)
    d.add_paragraph("Payment within 30 days.", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Price"
    t.cell(1, 0).text, t.cell(1, 1).text = "Service", "100"
    path = tmp_path / "contract.docx"
    d.save(str(path))
    return path


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2"])
    ws.append(["EMEA", 10, 12.0])
    ws.merge_cells("A4:C4")
    ws["A4"] = "merged note"
    path = tmp_path / "sales.xlsx"
    wb.save(str(path))
    return path
