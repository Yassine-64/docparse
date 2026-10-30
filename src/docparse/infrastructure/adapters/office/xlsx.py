"""XLSX via openpyxl: one Page per sheet, one table per sheet, merged cells filled."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from docparse.application.dto import ParseOptions
from docparse.domain.models import Block, BlockKind, Document, Page, Provenance, Table

SOURCE = "openpyxl"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fill_merged(ws: Worksheet) -> None:
    for rng in list(ws.merged_cells.ranges):
        value = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        for row in ws.iter_rows(rng.min_row, rng.max_row, rng.min_col, rng.max_col):
            for cell in row:
                cell.value = value


def _trim_columns(rows: list[list[str]]) -> list[list[str]]:
    width = max(len(r) for r in rows)
    last = max((i for i in range(width) if any(len(r) > i and r[i] for r in rows)), default=-1)
    return [r[: last + 1] for r in rows]


class XlsxParser:
    suffixes = frozenset({".xlsx", ".xlsm"})

    def __init__(self, max_rows: int = 5000) -> None:
        self.max_rows = max_rows

    async def parse(self, path: Path, options: ParseOptions) -> Document:
        return await asyncio.to_thread(self._parse, path)

    def _parse(self, path: Path) -> Document:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        pages = [self._sheet(ws, idx) for idx, ws in enumerate(wb.worksheets, start=1)]
        return Document(
            source_name=path.name,
            media_type=MEDIA_TYPE,
            pages=pages,
            metadata={"sheets": [ws.title for ws in wb.worksheets]},
        )

    def _sheet(self, ws: Worksheet, idx: int) -> Page:
        _fill_merged(ws)
        rows = [[_fmt(v) for v in r] for r in ws.iter_rows(values_only=True, max_row=self.max_rows)]
        rows = [r for r in rows if any(r)]
        prov = Provenance(page=idx, source=SOURCE)
        page = Page(
            number=idx,
            width=1.0,
            height=1.0,
            blocks=[Block(kind=BlockKind.HEADING, level=1, text=ws.title, provenance=prov)],
        )
        if rows:
            table = Table(rows=_trim_columns(rows))
            page.blocks.append(
                Block(kind=BlockKind.TABLE, text=table.to_markdown(), table=table, provenance=prov)
            )
        return page
